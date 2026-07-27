import csv
import io
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from django.db import transaction
from django.utils import timezone

from core.models import Agendamento, Paciente


TIPO_ATENDIMENTO_MAP = {
    "consulta": "consulta",
    "consultas": "consulta",
    "exame": "exame_auditivo",
    "exames": "exame_auditivo",
    "exame_auditivo": "exame_auditivo",
    "terapia": "terapia",
    "terapias": "terapia",
    "fonoaudiologia": "exame_auditivo",
    "fisioterapia": "terapia",
    "psicologia": "terapia",
    "terapia ocupacional": "terapia",
}

COLUNAS_ESPERADAS = [
    "nome_completo",
    "nome do paciente",
    "paciente",
    "nome",
    "cpf",
    "data_nascimento",
    "data de nascimento",
    "nascimento",
    "nome_mae",
    "nome da mae",
    "nome da mãe",
    "mae",
    "tipo_atendimento",
    "tipo de atendimento",
    "atendimento",
    "tipo",
    "servico",
    "serviço",
]


@dataclass
class LinhaImportada:
    nome_completo: str
    cpf: str
    data_nascimento: Optional[date] = None
    nome_mae: str = ""
    tipo_atendimento: str = "consulta"
    erros: list[str] = field(default_factory=list)
    valida: bool = True


@dataclass
class ResultadoImportacao:
    total: int = 0
    criados: int = 0
    ignorados: int = 0
    erros: list[dict] = field(default_factory=list)
    pacientes_criados: int = 0


def _apenas_digitos(valor: str) -> str:
    return re.sub(r"\D", "", valor)


def _normalizar_data(valor: Any) -> Optional[date]:
    if isinstance(valor, date):
        return valor
    if isinstance(valor, datetime):
        return valor.date()
    if not valor or not str(valor).strip():
        return None
    texto = str(valor).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    try:
        from pandas import Timestamp
        if isinstance(valor, Timestamp):
            return valor.date()
    except ImportError:
        pass
    return None


def _normalizar_tipo_atendimento(valor: Any) -> str:
    if not valor:
        return "consulta"
    texto = str(valor).strip().lower()
    return TIPO_ATENDIMENTO_MAP.get(texto, "consulta")


def _extrair_coluna(linha: dict, nomes_possiveis: list[str]) -> str:
    for nome in nomes_possiveis:
        if nome in linha:
            valor = linha[nome]
            if isinstance(valor, str):
                return valor.strip()
            return str(valor) if valor is not None else ""
    return ""


def _parse_linha_dict(linha: dict) -> LinhaImportada:
    nome = _extrair_coluna(linha, ["nome_completo", "nome do paciente", "paciente", "nome"])
    cpf = _apenas_digitos(_extrair_coluna(linha, ["cpf"]))
    data_nasc = _normalizar_data(_extrair_coluna(linha, ["data_nascimento", "data de nascimento", "nascimento"]))
    nome_mae = _extrair_coluna(linha, ["nome_mae", "nome da mae", "nome da mãe", "mae"])
    tipo = _normalizar_tipo_atendimento(_extrair_coluna(linha, ["tipo_atendimento", "tipo de atendimento", "atendimento", "tipo", "servico", "serviço"]))

    item = LinhaImportada(
        nome_completo=nome,
        cpf=cpf,
        data_nascimento=data_nasc,
        nome_mae=nome_mae,
        tipo_atendimento=tipo,
    )

    if not nome:
        item.erros.append("Nome do paciente é obrigatório")
        item.valida = False
    if not cpf or len(cpf) != 11:
        item.erros.append("CPF deve ter 11 dígitos")
        item.valida = False

    return item


def _ler_csv(arquivo: io.TextIOWrapper) -> list[LinhaImportada]:
    reader = csv.DictReader(arquivo)
    return [_parse_linha_dict(linha) for linha in reader]


def _ler_xlsx(caminho: str) -> list[LinhaImportada]:
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [str(c).strip().lower() if c else "" for c in next(linhas, [])]
    resultados = []
    for valores in linhas:
        linha_dict = dict(zip(cabecalho, valores))
        if any(v is not None for v in valores):
            resultados.append(_parse_linha_dict(linha_dict))
    wb.close()
    return resultados


def _ler_pdf(caminho: str) -> list[LinhaImportada]:
    import pdfplumber
    with pdfplumber.open(caminho) as pdf:
        texto_completo = "\n".join(page.extract_text() or "" for page in pdf.pages)

    linhas_texto = [l.strip() for l in texto_completo.split("\n") if l.strip()]
    resultados = []

    import csv as csv_module
    try:
        dialect = csv_module.Sniffer().sniff("\n".join(linhas_texto[:5]))
        reader = csv_module.DictReader(linhas_texto, dialect=dialect)
        return [_parse_linha_dict(linha) for linha in reader]
    except csv_module.Error:
        pass

    if not resultados:
        from io import StringIO
        try:
            reader = csv_module.DictReader(StringIO("\n".join(linhas_texto)))
            resultados = [_parse_linha_dict(linha) for linha in reader]
        except Exception:
            pass

    return resultados if resultados else []


def _ler_docx(caminho: str) -> list[LinhaImportada]:
    from docx import Document
    doc = Document(caminho)
    linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    import csv as csv_module
    from io import StringIO
    try:
        reader = csv_module.DictReader(StringIO("\n".join(linhas)))
        return [_parse_linha_dict(linha) for linha in reader]
    except csv_module.Error:
        pass

    tabelas = []
    for table in doc.tables:
        cabecalho = [cell.text.strip().lower() for cell in table.rows[0].cells]
        for row in table.rows[1:]:
            valores = [cell.text.strip() for cell in row.cells]
            tabelas.append(dict(zip(cabecalho, valores)))

    if tabelas:
        return [_parse_linha_dict(linha) for linha in tabelas]

    return []


def _extrair_data_filename(nome_arquivo: str) -> Optional[date]:
    match = re.search(r"(\d{4})[-_]?(\d{2})[-_]?(\d{2})", nome_arquivo)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            pass
    return None


def _salvar_temporario(arquivo) -> tuple[Path, str]:
    nome = getattr(arquivo, "name", "") or ""
    ext = Path(nome).suffix.lower() if nome else ""
    if not ext and hasattr(arquivo, "content_type"):
        ct = arquivo.content_type or ""
        ext_map = {
            "text/csv": ".csv",
            "application/pdf": ".pdf",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
            "application/vnd.ms-excel": ".xls",
            "application/msword": ".doc",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "image/png": ".png",
            "image/jpeg": ".jpg",
        }
        ext = ext_map.get(ct, ext)
    if not ext:
        ext = ".csv"
    fd, caminho = tempfile.mkstemp(suffix=ext)
    with os.fdopen(fd, "wb") as f:
        f.write(arquivo.read())
    return Path(caminho), nome


def process_appointment_file(
    arquivo, target_date: Optional[date] = None, senha_padrao: str = "123456"
) -> ResultadoImportacao:
    if target_date is None:
        target_date = timezone.localdate()

    caminho_tmp, nome_original = _salvar_temporario(arquivo)
    ext = caminho_tmp.suffix.lower()

    data_extraida = _extrair_data_filename(nome_original)
    if data_extraida and not target_date:
        target_date = data_extraida

    try:
        if ext in (".csv",):
            with caminho_tmp.open("r", encoding="utf-8-sig") as f:
                linhas = _ler_csv(f)
        elif ext in (".xlsx", ".xls"):
            linhas = _ler_xlsx(str(caminho_tmp))
        elif ext in (".pdf",):
            linhas = _ler_pdf(str(caminho_tmp))
        elif ext in (".doc", ".docx"):
            linhas = _ler_docx(str(caminho_tmp))
        elif ext in (".png", ".jpg", ".jpeg"):
            linhas = _ler_imagem(str(caminho_tmp))
        else:
            with caminho_tmp.open("r", encoding="utf-8-sig") as f:
                linhas = _ler_csv(f)
    finally:
        caminho_tmp.unlink(missing_ok=True)

    if not linhas:
        return ResultadoImportacao(erros=[{"linha": 0, "erros": ["Nenhum dado extraído do arquivo"]}])

    resultado = ResultadoImportacao(total=len(linhas))

    for idx, linha in enumerate(linhas, start=1):
        if not linha.valida:
            resultado.erros.append({"linha": idx, "nome": linha.nome_completo, "erros": linha.erros})
            continue

        with transaction.atomic():
            paciente, paciente_criado = Paciente.objects.get_or_create(
                cpf=linha.cpf,
                defaults={
                    "nome_completo": linha.nome_completo,
                    "data_nascimento": linha.data_nascimento,
                    "nome_mae": linha.nome_mae,
                },
            )
            if paciente_criado:
                paciente.set_senha(senha_padrao)
                paciente.save(update_fields=["senha_hash"])
                resultado.pacientes_criados += 1

            if Agendamento.objects.filter(
                paciente=paciente, data_agendamento=target_date
            ).exists():
                resultado.ignorados += 1
                continue

            Agendamento.objects.create(
                paciente=paciente,
                data_agendamento=target_date,
                tipo_atendimento=linha.tipo_atendimento,
            )
            resultado.criados += 1

    return resultado


def _ler_imagem(caminho: str) -> list[LinhaImportada]:
    try:
        import pytesseract
        from PIL import Image
        image = Image.open(caminho)
        texto = pytesseract.image_to_string(image, lang="por")
        linhas_texto = [l.strip() for l in texto.split("\n") if l.strip()]
        import csv as csv_module
        from io import StringIO
        reader = csv_module.DictReader(StringIO("\n".join(linhas_texto)))
        return [_parse_linha_dict(linha) for linha in reader]
    except ImportError:
        return []
    except Exception:
        return []
